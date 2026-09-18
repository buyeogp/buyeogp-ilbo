"""현행 일보 엑셀 5종을 판독해 M3 적재용 스테이징 CSV 를 만든다.

설계문서 §9 M3 원칙을 따른다 — **이벤트를 역산하지 않는다. 일별 스냅샷만 적재한다.**
엑셀이 계산해 둔 「당일두수」는 검증용으로만 싣고, 실제 당일두수는 DB 의 생성열이
다시 계산한다. 둘의 차이가 곧 설계문서 D3 「거의 매일 발생하는 계산 실수」의 실측치다.

판독하며 확인한 원본 구조 (README 참조)
  · 블록 앵커는 제목행이 아니라 **작성자 행**이다. 분만 8/15 는 제목이 C 열에 있어
    제목 기준으로 잡으면 8/14 블록에 붙어버린다
  · 분만사 포유자돈 유출은 N(포자전출) + O(이유전출) 두 열이다. 첫 블록만 라벨이 다르다
  · 육성사 「위탁판매(외부)」는 두수가 빠지지 않는다 — 소유권 이전이다

출력
  db/migration/m3_pen_daily.csv   스테이징 행
  db/migration/m3_report.md       판독·검산 품질 보고서

실행: python db/tools/parse_daily_reports.py
"""
import openpyxl, os, re, csv, io, datetime
from collections import Counter, defaultdict

BASE = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..'))
SRC = os.path.join(BASE, '자료들')
OUT = os.path.join(BASE, 'db', 'migration')
os.makedirs(OUT, exist_ok=True)

PEN = re.compile(r'^\s*(\d{1,2})\s*-\s*(\d{1,2})\s*$')
SUM_ROW = re.compile(r'^(총?합계|계)')
norm = lambda v: re.sub(r'\s+', '', '' if v is None else str(v))


def num(v):
    """두수 칸 판독. 빈칸·'/'·문자열은 0."""
    if v is None:
        return 0
    if isinstance(v, (int, float)):
        return int(round(v))
    s = str(v).strip()
    if s in ('', '/', '-', '.'):
        return 0
    return int(round(float(s))) if re.match(r'^-?\d+(\.\d+)?$', s) else 0


def raw(v):
    if v is None:
        return ''
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%Y-%m-%d')
    return str(v).strip().replace('\n', ' ')


def as_date(v):
    return v.strftime('%Y-%m-%d') if isinstance(v, (datetime.datetime, datetime.date)) else ''


def grid(ws, ncol):
    g = {}
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row,
                                         max_col=ncol, values_only=True), 1):
        g[r] = list(row)
    return g


def col(letter):
    c = 0
    for ch in letter:
        c = c * 26 + (ord(ch) - 64)
    return c


def at(g, r, letter):
    row = g.get(r)
    c = col(letter)
    return row[c - 1] if row and c <= len(row) else None


def find_blocks(g, ncol):
    """작성자 행을 블록 앵커로 삼는다. 전 파일에서 날짜와 1:1 로 대응한다."""
    anchors = []
    for r in sorted(g):
        row = g[r] or []
        if any(norm(v).startswith('작성자') for v in row[:4] if v is not None):
            d = ''
            for v in row[:ncol]:
                if isinstance(v, (datetime.datetime, datetime.date)):
                    d = v.strftime('%Y-%m-%d')
                    break
            anchors.append((r, d))
    out = []
    for i, (r, d) in enumerate(anchors):
        end = anchors[i + 1][0] - 1 if i + 1 < len(anchors) else max(g)
        out.append((r, end, d))
    return out


def data_rows(g, lo, hi, key_col):
    """블록 안에서 두수현황 표의 행 범위만 돌려준다 (구분 헤더 ~ 합계 앞)."""
    start = lo
    for r in range(lo, hi + 1):
        if norm(at(g, r, key_col)) == '구분':
            start = r + 1
            break
    for r in range(start, hi + 1):
        a = norm(at(g, r, key_col))
        if SUM_ROW.match(a) and a != '계':
            hi = r - 1
            break
    return range(start, hi + 1)


# ── 돈방 단위 돈사 ───────────────────────────────────────────────────
# owner_transfer : 두수가 빠지지 않는 소유권 이전 (육성사 위탁판매)
PEN_HOUSES = [
    ('자돈사,육성사.xlsx', ['자돈8월', '자돈9월'], 'JADON', {
        'pen': 'A', 'entry': 'C', 'sex': 'E', 'birth': 'F', 'weight': 'G',
        'opening': 'H', 'in': 'J', 'out': 'L', 'internal_out': 'N',
        'sold': None, 'owner_transfer': None, 'dead': 'P', 'closing': 'V', 'note': 'X',
    }, 30),
    ('자돈사,육성사.xlsx', ['육성8월', '육성9월'], 'YUKSUNG', {
        'pen': 'A', 'entry': 'C', 'sex': 'E', 'birth': 'F', 'weight': 'G',
        'opening': 'H', 'in': 'J', 'out': 'L', 'internal_out': 'N',
        'sold': None, 'owner_transfer': 'P', 'dead': 'R', 'closing': 'V', 'note': None,
    }, 30),
    ('검정사일보.xlsx', ['8월', '9월'], 'GEOMJUNG', {
        'pen': 'A', 'entry': 'C', 'sex': None, 'birth': 'E', 'weight': 'F',
        'opening': 'G', 'in': 'I', 'out': 'K', 'internal_out': 'M',
        'sold': 'O', 'owner_transfer': None, 'dead': 'Q', 'closing': 'U', 'note': None,
    }, 25),
    ('비육사일보.xlsx', ['8월', '9월'], 'BIYUK_M', {
        'pen': 'A', 'entry': 'C', 'sex': None, 'birth': 'E', 'weight': 'F',
        'opening': 'G', 'in': 'I', 'out': 'K', 'internal_out': 'M',
        'sold': 'O', 'owner_transfer': None, 'dead': 'Q', 'closing': 'U', 'note': None,
    }, 25),
]

JONGBU_HOUSES = {'순치사': 'SUNCHI', '종부사': 'JONGBU',
                 '임신1동': 'IMSIN1', '임신2동': 'IMSIN2'}
JONGBU_CATS = {
    '후보(수)': 'CAND_M', '후보(암)': 'CAND_F', '웅돈': 'BOAR',
    '이유모돈': 'WEAN_SOW', '이유돈': 'WEAN_PIG', '체류돈': 'STAY',
    '단기쳬류': 'STAY_S', '단기체류': 'STAY_S', '장기체류': 'STAY_L',
    '임신돈': 'PREG',
}

rows, diag, samples = [], defaultdict(Counter), defaultdict(list)


def emit(house, pen, cat, date, o, i, ot, io_, sold, dead, xclose, **kw):
    computed = o + i - ot - io_ - sold - dead
    diag[house]['행'] += 1
    if xclose is None:
        diag[house]['검산불가'] += 1
    elif computed == xclose:
        diag[house]['일치'] += 1
    else:
        diag[house]['불일치'] += 1
        if len(samples[house]) < 8:
            samples[house].append(
                f'{date} {pen or ""}{("/" + cat) if cat else ""}: 전일{o} +전입{i} '
                f'−전출{ot} −내부{io_} −판매{sold} −폐사{dead} = {computed} '
                f'/ 엑셀 {xclose} (차이 {xclose - computed:+d})')
    rec = {
        'house_code': house, 'pen_code': pen or '', 'category_code': cat or '',
        'report_date': date,
        'opening_head': o, 'in_head': i, 'out_head': ot,
        'internal_out_head': io_, 'sold_head': sold, 'dead_head': dead,
        'excel_closing': '' if xclose is None else xclose,
        'owner_transfer_head': 0, 'weaned_out_head': 0,
        'recurred_head': 0, 'aborted_head': 0, 'infertile_head': 0,
        'entry_date': '', 'birth_date_avg': '', 'entry_weight': '',
        'sex_mix': '', 'note': '',
    }
    rec.update(kw)
    rows.append(rec)


for fname, sheets, house, M, ncol in PEN_HOUSES:
    wb = openpyxl.load_workbook(os.path.join(SRC, fname), data_only=True, read_only=True)
    for sn in sheets:
        if sn not in wb.sheetnames:
            continue
        g = grid(wb[sn], ncol)
        for lo, hi, d in find_blocks(g, ncol):
            if not d:
                diag[house]['날짜없음블록'] += 1
                continue
            for r in data_rows(g, lo, hi, 'A'):
                m = PEN.match(norm(at(g, r, M['pen'])))
                if not m:
                    continue
                emit(house, f'{int(m.group(1))}-{int(m.group(2))}', None, d,
                     num(at(g, r, M['opening'])), num(at(g, r, M['in'])),
                     num(at(g, r, M['out'])), num(at(g, r, M['internal_out'])),
                     num(at(g, r, M['sold'])) if M['sold'] else 0,
                     num(at(g, r, M['dead'])), num(at(g, r, M['closing'])),
                     owner_transfer_head=(num(at(g, r, M['owner_transfer']))
                                          if M['owner_transfer'] else 0),
                     entry_date=raw(at(g, r, M['entry'])),
                     birth_date_avg=as_date(at(g, r, M['birth'])),
                     entry_weight=raw(at(g, r, M['weight'])),
                     sex_mix=raw(at(g, r, M['sex'])) if M['sex'] else '',
                     note=raw(at(g, r, M['note'])) if M['note'] else '')
    wb.close()

# ── 분만사 ───────────────────────────────────────────────────────────
# 분만복수(V) = 대기모돈 → 포유모돈 내부 이동.
# 포유자돈 유출 = N(포자전출) + O(이유전출).
wb = openpyxl.load_workbook(os.path.join(SRC, '분만사일보.xlsx'), data_only=True, read_only=True)
for sn in ['분만8월', '분만9월']:
    g = grid(wb[sn], 30)
    for lo, hi, d in find_blocks(g, 30):
        if not d:
            diag['BUNMAN']['날짜없음블록'] += 1
            continue
        for r in data_rows(g, lo, hi, 'A'):
            m = PEN.match(norm(at(g, r, 'A')))
            if not m:
                continue
            dong, no = int(m.group(1)), int(m.group(2))
            house = 'BUNMAN1' if dong == 1 else 'BUNMAN2'
            pen, entry = f'{dong}-{no}', raw(at(g, r, 'C'))
            litters = num(at(g, r, 'V'))
            weaned = num(at(g, r, 'O'))
            cats = [
                ('FARROW_WAIT', num(at(g, r, 'E')), num(at(g, r, 'I')), num(at(g, r, 'J')),
                 litters, num(at(g, r, 'P')), num(at(g, r, 'S')), 0),
                ('LACT_SOW', num(at(g, r, 'F')), num(at(g, r, 'K')) + litters,
                 num(at(g, r, 'L')), 0, num(at(g, r, 'Q')), num(at(g, r, 'T')), 0),
                ('SUCK', num(at(g, r, 'G')), num(at(g, r, 'M')),
                 num(at(g, r, 'N')) + weaned, 0, num(at(g, r, 'R')),
                 num(at(g, r, 'U')), weaned),
                ('WEANED', num(at(g, r, 'H')), 0, 0, 0, 0, None, 0),
            ]
            for cat, o, i, ot, io_, dead, xc, w in cats:
                if o == i == ot == io_ == dead == 0 and not xc:
                    continue
                emit(house, pen, cat, d, o, i, ot, io_, 0, dead, xc,
                     entry_date=entry, weaned_out_head=w)
wb.close()

# ── 종부사 ───────────────────────────────────────────────────────────
# 사고내역(재발·유산·불임)은 축종 간 이동이므로 내부전출로 본다.
wb = openpyxl.load_workbook(os.path.join(SRC, '종부사일보.xlsx'), data_only=True, read_only=True)
for sn in ['종부8월', '종부9월']:
    g = grid(wb[sn], 25)
    for lo, hi, d in find_blocks(g, 25):
        if not d:
            diag['JONGBU']['날짜없음블록'] += 1
            continue
        cur = None
        for r in range(lo, hi + 1):
            b, c = norm(at(g, r, 'B')), norm(at(g, r, 'C'))
            if b in JONGBU_HOUSES:
                cur = JONGBU_HOUSES[b]
            if b == '합계':
                break
            if not cur or c in ('', '계'):
                continue
            cat = JONGBU_CATS.get(c)
            if not cat:
                diag['JONGBU'][f'미매핑:{c}'] += 1
                continue
            rec, ab, inf = (num(at(g, r, 'J')), num(at(g, r, 'K')), num(at(g, r, 'L')))
            emit(cur, None, cat, d,
                 num(at(g, r, 'D')), num(at(g, r, 'F')), num(at(g, r, 'H')),
                 rec + ab + inf, num(at(g, r, 'M')), num(at(g, r, 'N')),
                 num(at(g, r, 'P')),
                 recurred_head=rec, aborted_head=ab, infertile_head=inf)
wb.close()

# ── 중복 키 검사 ─────────────────────────────────────────────────────
keyed = Counter((r['house_code'], r['pen_code'], r['category_code'], r['report_date'])
                for r in rows)
dups = {k: v for k, v in keyed.items() if v > 1}

# ── 전일두수 연속성 (설계문서 §1.5) ──────────────────────────────────
by_key = defaultdict(dict)
for r in rows:
    by_key[(r['house_code'], r['pen_code'], r['category_code'])][r['report_date']] = r

cont_ok = cont_bad = cont_na = 0
cont_samples = []
for key, days in by_key.items():
    ds = sorted(days)
    for a, b in zip(ds, ds[1:]):
        prev, cur = days[a], days[b]
        if prev['excel_closing'] == '':
            cont_na += 1
            continue
        if int(prev['excel_closing']) == cur['opening_head']:
            cont_ok += 1
        else:
            cont_bad += 1
            if len(cont_samples) < 10:
                label = f'{key[0]} {key[1] or ""}{("/" + key[2]) if key[2] else ""}'
                cont_samples.append(
                    f'{label:26s} {a} 당일 {prev["excel_closing"]:>4} '
                    f'→ {b} 전일 {cur["opening_head"]:>4}')

# ── 출력 ─────────────────────────────────────────────────────────────
cols = ['house_code', 'pen_code', 'category_code', 'report_date',
        'opening_head', 'in_head', 'out_head', 'internal_out_head',
        'sold_head', 'dead_head', 'excel_closing',
        'owner_transfer_head', 'weaned_out_head',
        'recurred_head', 'aborted_head', 'infertile_head',
        'entry_date', 'birth_date_avg', 'entry_weight', 'sex_mix', 'note']
rows.sort(key=lambda r: (r['report_date'], r['house_code'], r['pen_code'], r['category_code']))
csv_path = os.path.join(OUT, 'm3_pen_daily.csv')
with io.open(csv_path, 'w', encoding='utf-8', newline='') as f:
    w = csv.DictWriter(f, fieldnames=cols)
    w.writeheader()
    w.writerows(rows)

dates = sorted({r['report_date'] for r in rows})
tot = sum(diag[h]['행'] for h in diag)
match = sum(diag[h]['일치'] for h in diag)
bad = sum(diag[h]['불일치'] for h in diag)
na = sum(diag[h]['검산불가'] for h in diag)
den = match + bad

L = []
w = L.append
w('# M3 과거 데이터 판독 보고서')
w('')
w('현행 일보 엑셀 5종을 판독해 `m3_pen_daily.csv` 를 만들었다.')
w('설계문서 §9 M3 원칙대로 **이벤트를 역산하지 않고 일별 스냅샷만** 담았다.')
w('')
w(f'- 생성: `db/tools/parse_daily_reports.py`')
w(f'- 행 수: **{tot:,}**   ·   일자: **{dates[0]} ~ {dates[-1]}** ({len(dates)}일)')
w(f'- 중복 키: **{len(dups)}건**')
w('')
w('## 검산 결과')
w('')
w('`전일두수 + 전입 − 전출 − 내부전출 − 판매 − 폐사` 를 엑셀의 「당일두수」와 대조했다.')
w('')
w('| 돈사 | 행 | 일치 | 불일치 | 검산불가 | 일치율 |')
w('|---|---:|---:|---:|---:|---:|')
for h in sorted(diag):
    c = diag[h]
    dd = c['일치'] + c['불일치']
    rate = f'{100.0 * c["일치"] / dd:.1f}%' if dd else '—'
    w(f'| {h} | {c["행"]:,} | {c["일치"]:,} | {c["불일치"]:,} | {c["검산불가"]:,} | {rate} |')
if den:
    w(f'| **계** | **{tot:,}** | **{match:,}** | **{bad:,}** | **{na:,}** | '
      f'**{100.0 * match / den:.2f}%** |')
w('')
if bad:
    w('### 불일치 표본')
    w('')
    w('```')
    for h in sorted(samples):
        for s in samples[h][:5]:
            w(f'{h:10s} {s}')
    w('```')
    w('')
w('## 전일두수 연속성')
w('')
w('어제의 「당일두수」가 오늘의 「전일두수」로 그대로 넘어왔는지 전수 검사했다.')
w('')
w(f'- 검사 쌍 **{cont_ok + cont_bad:,}** · 일치 **{cont_ok:,}** · 불일치 **{cont_bad:,}** '
  f'· 대조불가 {cont_na:,}')
w('')
if cont_samples:
    w('```')
    for s in cont_samples:
        w(s)
    w('```')
    w('')
w('## 판독하며 확인한 원본 구조')
w('')
w('설계문서에 없거나 다르게 적힌 것들이다.')
w('')
w('- **블록 앵커는 작성자 행이다.** 분만 8/15 블록은 제목이 A 열이 아니라 C 열에 있어')
w('  제목 기준으로 잡으면 8/14 블록에 통째로 붙는다. 설계문서 §11.3 이 「분만사 8/15 결측」')
w('  으로 기록한 건은 **결측이 아니라 제목행 위치 이상**이며 데이터는 존재한다')
w('- **분만사 포유자돈 유출은 두 열이다** — N(포자전출) + O(이유전출).')
w('  28개 블록 중 27개가 이 배치고, 첫 블록(8/6)만 라벨이 `M=전입 / O=전출` 로 다르다 (D6)')
w('- **분만복수(V)는 대기모돈 → 포유모돈 내부 이동이다.** 별도 이동 열이 없고 이 값으로만 잡힌다')
w('- **육성사 「위탁판매(외부)」는 두수가 빠지지 않는다.** 소유권만 중앙축산으로 넘어가고')
w('  돼지는 육성사에 그대로 있다 (설계문서 §4.10 「소유권 이전: 자돈 75일령, 숫퇘지에 한함」).')
w('  `owner_transfer_head` 로 따로 실었다 — 두수 등식에서 빼면 안 된다')
w('- **분만사 「이유자돈」 칸은 전일두수만 있고 당일두수 열이 없다.** 검산 대상에서 제외했다')
w('')
w('## 적재 방법')
w('')
w('```bash')
w('psql -v ON_ERROR_STOP=1 -f db/migration/m3_00_staging.sql')
w("psql -c \"\\copy app.m3_daily_raw FROM 'db/migration/m3_pen_daily.csv' CSV HEADER\"")
w('psql -v ON_ERROR_STOP=1 -f db/migration/m3_01_load.sql')
w('```')

io.open(os.path.join(OUT, 'm3_report.md'), 'w', encoding='utf-8').write('\n'.join(L) + '\n')

print(f'행 {tot:,} / 일자 {len(dates)}일 ({dates[0]} ~ {dates[-1]}) / 중복키 {len(dups)}')
print(f'검산 일치 {match:,} · 불일치 {bad:,} · 불가 {na:,}'
      + (f' → {100.0 * match / den:.2f}%' if den else ''))
print(f'연속성 일치 {cont_ok:,} · 불일치 {cont_bad:,}')
for h in sorted(diag):
    extra = {k: v for k, v in diag[h].items()
             if k not in ('행', '일치', '불일치', '검산불가')}
    print(f'  {h:10s} 행 {diag[h]["행"]:5,}  불일치 {diag[h]["불일치"]:3,}'
          + (f'  {extra}' if extra else ''))
print('→', csv_path)
