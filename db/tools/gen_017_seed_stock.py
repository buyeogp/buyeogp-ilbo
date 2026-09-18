"""재고관리 엑셀 2종에서 약품·사료 마스터를 뽑아 017_seed_medicine_feed.sql 을 만든다.

설계문서 §4.9.2 는 「재고 파일에는 휴약기간이 12개 품목에만 있고, 거래명세표에는
대부분 인쇄되어 있다」고 한다. 그래서 두 원천을 합친다.
  · 품목 목록·제형·단가·유효기간 → 202608 재고관리(약품).xlsx
  · 휴약기간                      → AG동물약품 거래명세표 (아래 WITHDRAWAL 상수)

사료는 202608 재고관리(사료).xlsx 에서 품목·월별 단가·일자별 투입량을 뽑는다.

출력
  db/017_seed_medicine_feed.sql      약품·사료 마스터 + 단가 이력
  db/migration/m3_feed_delivery.csv  일자별 사료 투입 (M3 확장)
  db/migration/m3_feed_report.md     판독 보고서

실행: python db/tools/gen_017_seed_stock.py
"""
import openpyxl, os, re, io, csv, datetime
from collections import OrderedDict, defaultdict, Counter

BASE = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..'))
SRC = os.path.join(BASE, '자료들')
OUT_SQL = os.path.join(BASE, 'db', '017_seed_medicine_feed.sql')
OUT_MIG = os.path.join(BASE, 'db', 'migration')
os.makedirs(OUT_MIG, exist_ok=True)

norm = lambda v: re.sub(r'\s+', '', '' if v is None else str(v))
q = lambda s: "'" + str(s).replace("'", "''") + "'" if s not in (None, '') else 'NULL'

# 거래명세표 6매에서 확인된 휴약기간 (설계문서 §4.9.2 표)
WITHDRAWAL = {
    '안티펜SM': 35, '툴라신': 33, '덱소론': 28, '덱사메타손100': 28,
    '복합부스코판': 28, '베트리목신': 21, '네오신첨가제': 21, '부타D PPS': 20,
    '빌아목사펜': 15, '바이트릴100주': 10, '슈라목스수용산': 10,
    '네오신M수용산': 5, 'HD피린': 5, '세프론-주': 1, '엑스포-피': 1,
    '루텔라이스': 1, '고나돈': 1, '비에스 바시트렉스100': 0, '옥시토신': 0,
}
# 재고파일 「구분」 → (약효분류 category, 제형 form)
KIND = {
    '주사제': (None, '주사제'), '첨가제': (None, '첨가제'), '수용산': (None, '수용산'),
    '백신': ('백신', '백신'), '항생제': ('항생제', None), '소독': ('소독', '소독제'),
    '소독제': ('소독', '소독제'), '기타': ('기타', None), '외품': ('외품', None),
    '호르몬제': ('호르몬제', None), '소염제': ('소염제', None), '해열제': ('해열제', None),
}

# ── 약품 ─────────────────────────────────────────────────────────────
med = OrderedDict()
kinds_seen = Counter()
wb = openpyxl.load_workbook(os.path.join(SRC, '202608 재고관리(약품).xlsx'),
                            data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
for row in ws.iter_rows(min_row=5, max_row=ws.max_row, max_col=9, values_only=True):
    supplier, name, wd, exp, kind, unit, price = (row[1], row[2], row[3], row[4],
                                                  row[5], row[6], row[7])
    name = norm(name)
    if not name or name in ('품목',):
        continue
    rx = name.startswith('(처)')
    clean = re.sub(r'^\(처\)', '', name).strip()
    # 제조사가 이름 끝 괄호에 붙는 경우: 슈라목스수용산(버박)
    maker = None
    m = re.match(r'^(.*?)\(([^()]{1,8})\)$', clean)
    if m and not re.search(r'\d', m.group(2)):
        clean, maker = m.group(1).strip(), m.group(2)
    kind = norm(kind)
    kinds_seen[kind] += 1
    cat, form = KIND.get(kind, ('기타', None))
    spec = norm(unit) or None
    key = (clean, spec)
    if key in med:
        continue
    med[key] = {
        'name': clean, 'maker': maker, 'rx': rx,
        'withdrawal': WITHDRAWAL.get(clean, int(wd) if isinstance(wd, (int, float)) else 0),
        'wd_src': 'invoice' if clean in WITHDRAWAL else ('stock' if isinstance(wd, (int, float)) else 'none'),
        'category': cat or '기타', 'form': form, 'spec': spec,
        'price': int(price) if isinstance(price, (int, float)) else None,
        'expiry': exp.strftime('%Y-%m-%d') if isinstance(exp, (datetime.datetime, datetime.date)) else None,
        'supplier': norm(supplier) or 'AG',
    }
wb.close()

# ── 사료 ─────────────────────────────────────────────────────────────
HOUSE_MAP = {
    '순치사': 'SUNCHI', '종부사': 'JONGBU', '임신1동': 'IMSIN1', '임신2동': 'IMSIN2',
    '분만사1동': 'BUNMAN1', '분만사2동': 'BUNMAN2',
    '초기자돈1동': 'JADON', '초기자돈2동': 'JADON',
    '육성1(수)': 'YUKSUNG', '육성2(암)': 'YUKSUNG',
    '검정사1(암)': 'GEOMJUNG', '검정사2(암)': 'GEOMJUNG',
    '검정사3(암)': 'GEOMJUNG', '검정사4(암)': 'GEOMJUNG',
    '비육1(수)': 'BIYUK_M', '비육2(암)': 'BIYUK_F',
    '계류장': 'GYERYU',
    '욱성2(암)': 'YUKSUNG',          # 원본 오타 (육성2(암))
}
SHEET_MONTH = re.compile(r'^(\d{2})(\d{2})$')

feed_items = OrderedDict()
price_hist = {}                      # (품목, 'YYYY-MM-01') -> 단가
delivery = defaultdict(int)          # (house, item, date) -> kg
house_items = defaultdict(Counter)   # 사료 구분별 품목 사용 (P-1 단서)
unmapped = Counter()

wb = openpyxl.load_workbook(os.path.join(SRC, '202608 재고관리(사료).xlsx'),
                            data_only=True, read_only=True)
for sn in wb.sheetnames:
    m = SHEET_MONTH.match(sn)
    if not m:
        continue
    yy, mm = 2000 + int(m.group(1)), int(m.group(2))
    ym = f'{yy:04d}-{mm:02d}-01'
    ws = wb[sn]
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, max_col=40, values_only=True):
        house_raw, item, price = norm(row[0]), norm(row[1]), row[2]
        if house_raw == '합계':
            break                     # 아래는 품목별 요약 블록이다 — 돈사 표가 아니다
        if house_raw.startswith('※') or house_raw in ('', '구분'):
            continue
        if house_raw not in HOUSE_MAP:
            if item:
                unmapped[house_raw] += 1
            continue
        if not item or not isinstance(price, (int, float)):
            continue
        feed_items.setdefault(item, {'name': item})
        house_items[house_raw][item] += 1
        prev = price_hist.get((item, ym))
        if prev is None or price < prev:
            price_hist[(item, ym)] = price
        hc = HOUSE_MAP[house_raw]
        for day in range(1, 32):
            v = row[2 + day] if 2 + day < len(row) else None
            if isinstance(v, (int, float)) and v:
                try:
                    d = datetime.date(yy, mm, day).isoformat()
                except ValueError:
                    continue
                delivery[(hc, item, d)] += int(round(v))
wb.close()

for it in feed_items:
    feed_items[it]['pack'] = ('지대' if '지대' in it else '벌크' if '벌크' in it else None)

# ── SQL 생성 ─────────────────────────────────────────────────────────
L = []
w = L.append
w('-- =====================================================================')
w('-- 017 약품 · 사료 마스터 — 재고관리 엑셀 판독본')
w('--')
w('-- 015 의 약품 20품목(거래명세표)·사료 8품목(추정)을 대체한다.')
w('-- 원천 : 202608 재고관리(약품).xlsx  — 품목·제형·규격·단가·유효기간')
w('--        AG동물약품 거래명세표 6매   — 휴약기간 (설계문서 §4.9.2)')
w('--        202608 재고관리(사료).xlsx  — 품목·월별 단가')
w('-- 생성 : db/tools/gen_017_seed_stock.py')
w('-- =====================================================================')
w('SET search_path = app, sec, extensions, public;')
w('')
w(f'-- ── 약품 {len(med)}품목 ─────────────────────────────────────────────')
w('INSERT INTO medicine (name, maker, is_prescription, withdrawal_days,')
w('                      category, form, spec, unit, unit_price, expiry_date, supplier_id)')
w('SELECT v.name, v.maker, v.rx, v.wd, v.cat::medicine_category,')
w('       v.form::medicine_form, v.spec, COALESCE(v.spec, \'EA\'), v.price, v.expiry::date, s.id')
w('  FROM supplier s, (VALUES')
body = []
for k, d in med.items():
    body.append('    ({n}, {mk}, {rx}, {wd}, {cat}, {fm}, {sp}, {pr}, {ex})'.format(
        n=q(d['name']), mk=q(d['maker']), rx='true' if d['rx'] else 'false',
        wd=d['withdrawal'], cat=q(d['category']),
        fm=q(d['form']) if d['form'] else 'NULL', sp=q(d['spec']),
        pr=d['price'] if d['price'] is not None else 'NULL',
        ex=q(d['expiry']) if d['expiry'] else 'NULL'))
w(',\n'.join(body))
w('  ) AS v(name, maker, rx, wd, cat, form, spec, price, expiry)')
w(" WHERE s.name = '(주)AG동물약품'")
w('ON CONFLICT (name, spec) DO NOTHING;')
w('')
w('-- 구제역 백신 — 축협, 국가 관리이므로 휴약기간 없음 (R-E)')
w('INSERT INTO medicine (name, is_prescription, withdrawal_days, category, unit, supplier_id)')
w("SELECT '구제역 백신', false, 0, '백신', 'EA', id FROM supplier WHERE name = '축협'")
w('ON CONFLICT (name, spec) DO NOTHING;')
w('')
w(f'-- ── 사료 {len(feed_items)}품목 ─────────────────────────────────────────────')
w('INSERT INTO feed (name, pack) VALUES')
w(',\n'.join(f'  ({q(d["name"])}, {q(d["pack"]) if d["pack"] else "NULL"})'
             for d in feed_items.values()))
w('ON CONFLICT (name) DO NOTHING;')
w('')
months = sorted({ym for _, ym in price_hist})
w(f'-- ── 사료 단가 이력 {len(price_hist)}건 / {len(months)}개월 ({months[0][:7]} ~ {months[-1][:7]}) ──')
w('INSERT INTO feed_price_history (feed_id, valid_from, unit_price)')
w('SELECT f.id, v.vf::date, v.price')
w('  FROM (VALUES')
w(',\n'.join(f'    ({q(it)}, {q(ym)}, {price})'
             for (it, ym), price in sorted(price_hist.items(), key=lambda x: (x[0][1], x[0][0]))))
w('  ) AS v(name, vf, price)')
w('  JOIN feed f ON f.name = v.name')
w('ON CONFLICT (feed_id, valid_from) DO NOTHING;')
io.open(OUT_SQL, 'w', encoding='utf-8').write('\n'.join(L) + '\n')

# ── 사료 투입 CSV ────────────────────────────────────────────────────
dpath = os.path.join(OUT_MIG, 'm3_feed_delivery.csv')
with io.open(dpath, 'w', encoding='utf-8', newline='') as f:
    wtr = csv.writer(f)
    wtr.writerow(['house_code', 'feed_name', 'event_date', 'qty_kg'])
    for (h, it, d), kg in sorted(delivery.items(), key=lambda x: (x[0][2], x[0][0], x[0][1])):
        wtr.writerow([h, it, d, kg])

# ── 보고서 ───────────────────────────────────────────────────────────
ddates = sorted({d for _, _, d in delivery})
R = []
w = R.append
w('# 재고관리 엑셀 판독 보고서')
w('')
w('## 약품')
w('')
w(f'- 품목 **{len(med)}개** (설계문서 §4.9.2 의 거래명세표 20품목보다 넓다)')
src = Counter(d['wd_src'] for d in med.values())
w(f'- 휴약기간 출처: 거래명세표 {src["invoice"]}건 · 재고파일 {src["stock"]}건 · '
  f'**미상 {src["none"]}건**')
w('')
w('휴약기간이 미상인 품목은 0 으로 넣었다. V9(휴약기간 미경과 출하 차단)가 이 값을')
w('그대로 쓰므로, **실제 휴약기간이 있는데 0 으로 남으면 차단이 작동하지 않는다.**')
w('거래명세표를 더 받거나 품목별로 확인해 채워야 한다.')
w('')
w('| 휴약기간 미상 품목 | 규격 | 제형 |')
w('|---|---|---|')
for d in med.values():
    if d['wd_src'] == 'none' and d['rx']:
        w(f'| {d["name"]} | {d["spec"] or "—"} | {d["form"] or "—"} |')
w('')
w(f'재고파일 「구분」 값: {dict(kinds_seen)}')
w('')
w('## 사료')
w('')
w(f'- 품목 **{len(feed_items)}개**: {", ".join(feed_items)}')
w(f'- 단가 이력 **{len(price_hist)}건** / {len(months)}개월 ({months[0][:7]} ~ {months[-1][:7]})')
w(f'- 일자별 투입 **{len(delivery):,}건** ({ddates[0]} ~ {ddates[-1]})')
w('')
w('설계문서 §4.8 대로 이 투입 기록이 사실상 급이량 원장이다.')
w('돈군별 두수·증체와 결합하면 FCR 이 추가 입력 없이 나온다.')
w('')
w('## 사료 파일의 돈사 구분이 일보보다 세분화되어 있다')
w('')
w('일보는 돈사 12개지만 사료 파일은 17개로 나눈다.')
w('')
w('| 사료 파일 | 일보 돈사 |')
w('|---|---|')
for k, v in HOUSE_MAP.items():
    w(f'| {k} | {v} |')
w('')
w('**P-1 단서.** 사료 파일은 `비육1(수)` 와 `비육2(암)` 을 따로 관리한다 —')
w('비육사(암)이 실재하는 별도 단위라는 뜻이다. 또 `검정사1(암)`~`검정사4(암)` 4개와')
w('`비육2(암)` 1개를 합치면 5개인데, 검정사일보의 동 번호도 1~5동이다.')
w('검정사일보의 5개 동 중 하나가 비육사(암)일 가능성이 있다.')
w('')
w('다만 설계문서 §4.3 은 「검정사 5-1~5-4 수령 400두」라고 적어 5동을 검정사로 본다.')
w('둘 중 하나가 틀렸으므로 현장에 물어야 한다 — **「검정사일보의 1~5동 중 비육사(암)은')
w('어느 동입니까?」** 가 정확한 질문이다.')
w('')
w('### 사료 구분별 사용 품목')
w('')
w('```')
for h in HOUSE_MAP:
    if h in house_items:
        w(f'{h:14s} {", ".join(sorted(house_items[h]))}')
w('```')
if unmapped:
    w('')
    w(f'미매핑 구분: {dict(unmapped)}')

io.open(os.path.join(OUT_MIG, 'm3_feed_report.md'), 'w', encoding='utf-8').write('\n'.join(R) + '\n')

print(f'약품 {len(med)}품목 (휴약 출처: {dict(src)})')
print(f'사료 {len(feed_items)}품목 · 단가이력 {len(price_hist)}건 · 투입 {len(delivery):,}건')
print(f'투입 기간 {ddates[0]} ~ {ddates[-1]}')
if unmapped:
    print('미매핑 사료 구분:', dict(unmapped))
print('→', OUT_SQL)
print('→', dpath)
